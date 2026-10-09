"""Tiny document generators for tests (no binary fixtures in the repository)."""

from __future__ import annotations

import io

import docx


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(page_texts: list[str | None]) -> bytes:
    """Build a valid PDF with one page per entry (``None`` = page without text)."""
    objects: dict[int, bytes] = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    kids: list[str] = []
    next_id = 4
    for text in page_texts:
        page_id, content_id = next_id, next_id + 1
        next_id += 2
        kids.append(f"{page_id} 0 R")
        stream = (
            b""
            if text is None
            else f"BT /F1 12 Tf 72 720 Td ({_pdf_escape(text)}) Tj ET".encode("latin-1")
        )
        objects[page_id] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>"
        ).encode()
        objects[content_id] = (
            f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream"
        )
    objects[2] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>".encode()

    out = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for object_id in sorted(objects):
        offsets[object_id] = len(out)
        out += f"{object_id} 0 obj\n".encode() + objects[object_id] + b"\nendobj\n"
    xref_at = len(out)
    size = max(objects) + 1
    out += f"xref\n0 {size}\n".encode() + b"0000000000 65535 f \n"
    for object_id in range(1, size):
        out += f"{offsets[object_id]:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {size} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode()
    return bytes(out)


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()

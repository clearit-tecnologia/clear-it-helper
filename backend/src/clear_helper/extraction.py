"""Baseline text extraction (pipeline ``baseline-v1``), one entry per page.

PDF pages are the real pages; every other format is a single page (page 1). This is the
deliberately simple baseline: no layout, no OCR, no structure (those arrive in Phase 2).
All functions are synchronous and CPU-bound; callers run them in a thread.
"""

from __future__ import annotations

import io
import logging
import re
import zipfile
from dataclasses import dataclass
from enum import StrEnum
from pathlib import PurePosixPath

import docx
from bs4 import BeautifulSoup
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader
from pypdf.errors import PdfReadError

logger = logging.getLogger(__name__)

NO_TEXT_PDF_MESSAGE = "Documento sem texto extraível (OCR chega na Fase 2)"
NO_TEXT_MESSAGE = "Documento sem texto extraível"


class DocumentFormat(StrEnum):
    PDF = "pdf"
    DOCX = "docx"
    TXT = "txt"
    MARKDOWN = "markdown"
    HTML = "html"


CONTENT_TYPES: dict[DocumentFormat, str] = {
    DocumentFormat.PDF: "application/pdf",
    DocumentFormat.DOCX: (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ),
    DocumentFormat.TXT: "text/plain",
    DocumentFormat.MARKDOWN: "text/markdown",
    DocumentFormat.HTML: "text/html",
}

_EXTENSIONS: dict[str, DocumentFormat] = {
    ".pdf": DocumentFormat.PDF,
    ".docx": DocumentFormat.DOCX,
    ".txt": DocumentFormat.TXT,
    ".text": DocumentFormat.TXT,
    ".md": DocumentFormat.MARKDOWN,
    ".markdown": DocumentFormat.MARKDOWN,
    ".html": DocumentFormat.HTML,
    ".htm": DocumentFormat.HTML,
}

_CONTENT_TYPE_FORMATS: dict[str, DocumentFormat] = {
    **{value: key for key, value in CONTENT_TYPES.items()},
    "text/x-markdown": DocumentFormat.MARKDOWN,
    "application/xhtml+xml": DocumentFormat.HTML,
}

_TEXT_FORMATS = frozenset({DocumentFormat.TXT, DocumentFormat.MARKDOWN, DocumentFormat.HTML})
_INVISIBLE_HTML_TAGS = ("script", "style", "noscript", "template", "head", "svg", "iframe")
_BLANK_LINES = re.compile(r"\n{3,}")
_TRAILING_SPACES = re.compile(r"[ \t]+\n")


class ExtractionError(Exception):
    """Extraction failed; ``str(exc)`` is a short pt-BR message safe to show to the user."""


@dataclass(frozen=True, slots=True)
class PageText:
    page: int
    text: str


@dataclass(frozen=True, slots=True)
class ExtractedDocument:
    page_count: int
    pages: list[PageText]


def detect_format(filename: str, content_type: str | None, head: bytes) -> DocumentFormat | None:
    """Return the document format, or ``None`` when it is not supported.

    The extension wins over the declared content type (browsers often send
    ``application/octet-stream``); binary formats are confirmed by their magic bytes.
    """
    suffix = PurePosixPath(filename.replace("\\", "/")).suffix.lower()
    fmt = _EXTENSIONS.get(suffix)
    if fmt is None and content_type:
        fmt = _CONTENT_TYPE_FORMATS.get(content_type.split(";", 1)[0].strip().lower())
    if fmt is None:
        return None
    if fmt is DocumentFormat.PDF and not head.lstrip()[:5].startswith(b"%PDF-"):
        return None
    if fmt is DocumentFormat.DOCX and not head.startswith(b"PK"):
        return None
    if fmt in _TEXT_FORMATS and b"\x00" in head[:4096] and not _looks_like_utf16(head):
        return None
    return fmt


def _looks_like_utf16(data: bytes) -> bool:
    return data.startswith((b"\xff\xfe", b"\xfe\xff"))


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = _TRAILING_SPACES.sub("\n", text)
    text = _BLANK_LINES.sub("\n\n", text)
    return text.strip()


def decode_text(data: bytes) -> str:
    if _looks_like_utf16(data):
        return data.decode("utf-16")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _extract_pdf(data: bytes) -> ExtractedDocument:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            # Many PDFs are "encrypted" with an empty user password (only permissions set).
            try:
                reader.decrypt("")
            except Exception as exc:
                raise ExtractionError("PDF protegido por senha") from exc
        pages: list[PageText] = []
        for number, page in enumerate(reader.pages, start=1):
            pages.append(PageText(page=number, text=normalize_text(page.extract_text() or "")))
    except ExtractionError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as exc:
        raise ExtractionError("PDF inválido ou corrompido") from exc
    if not any(page.text for page in pages):
        raise ExtractionError(NO_TEXT_PDF_MESSAGE)
    return ExtractedDocument(page_count=len(pages), pages=[p for p in pages if p.text])


def _table_text(table: Table) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _extract_docx(data: bytes) -> str:
    try:
        document = docx.Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ExtractionError("DOCX inválido ou corrompido") from exc
    blocks: list[str] = []
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            blocks.append(item.text)
        elif isinstance(item, Table):
            blocks.append(_table_text(item))
    return "\n".join(blocks)


def _extract_html(data: bytes) -> str:
    soup = BeautifulSoup(data, "html.parser")
    for tag in soup.find_all(_INVISIBLE_HTML_TAGS):
        tag.decompose()
    return soup.get_text("\n")


def extract(data: bytes, fmt: DocumentFormat) -> ExtractedDocument:
    """Extract the text of ``data`` page by page; raises ``ExtractionError`` on failure."""
    if fmt is DocumentFormat.PDF:
        return _extract_pdf(data)
    if fmt is DocumentFormat.DOCX:
        text = _extract_docx(data)
    elif fmt is DocumentFormat.HTML:
        text = _extract_html(data)
    else:
        text = decode_text(data)
    text = normalize_text(text)
    if not text:
        raise ExtractionError(NO_TEXT_MESSAGE)
    return ExtractedDocument(page_count=1, pages=[PageText(page=1, text=text)])
